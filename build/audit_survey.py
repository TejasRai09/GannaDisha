# -*- coding: utf-8 -*-
"""Column-by-column audit of the raw plot survey.

    python build/audit_survey.py  ->  build/cache/audit_2627.json

Reads the workbook once, streaming, and records for every one of the 66
columns: how often it is blank, how often it carries a placeholder that means
"not recorded" without being blank, and what values it actually holds.

The placeholder part is the reason this exists. A column that is 100% filled
can still be empty in every sense that matters - the ERP writes 0 for an
unpicked dropdown, and "NONE" for a disease nobody looked for. Counting nulls
alone reports those columns as complete.
"""

import json
import os
import sys
from collections import Counter

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.environ.get("SURVEY_XLSX") or os.path.join(
    ROOT, "Plot Wise Survey Report 2026-2027.xlsx"
)
OUT = os.path.join(HERE, "cache", "audit_2627.json")

# Values that are present but carry no information.
BLANKISH = {"", "-", "--", "NA", "N/A", "NULL", "NONE", "NIL", "0", "0.0", "?"}


def main() -> None:
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(rows))]
    n = len(header)

    blank = [0] * n           # None or empty string
    zeroish = [0] * n         # present but meaningless
    counts = [Counter() for _ in range(n)]
    numeric_ok = [0] * n
    numeric_zero = [0] * n
    nmin = [None] * n
    nmax = [None] * n
    total = 0

    for row in rows:
        total += 1
        if total % 25000 == 0:
            print(f"  {total:,} rows", file=sys.stderr)
        for i in range(n):
            v = row[i] if i < len(row) else None

            if v is None:
                blank[i] += 1
                continue
            s = str(v).strip()
            if s == "":
                blank[i] += 1
                continue
            if s.upper() in BLANKISH:
                zeroish[i] += 1

            # Keep the value table bounded - a grower-name column would
            # otherwise hold 74,000 entries and tell us nothing new.
            if len(counts[i]) < 3000:
                counts[i][s] += 1

            try:
                f = float(s)
            except ValueError:
                continue
            numeric_ok[i] += 1
            if f == 0:
                numeric_zero[i] += 1
            nmin[i] = f if nmin[i] is None else min(nmin[i], f)
            nmax[i] = f if nmax[i] is None else max(nmax[i], f)

    wb.close()

    out = {"file": os.path.basename(SRC), "rows": total, "columns": []}
    for i, name in enumerate(header):
        c = counts[i]
        top = c.most_common(8)
        distinct = len(c)
        out["columns"].append({
            "name": name,
            "blank": blank[i],
            "blankPct": round(blank[i] / total * 100, 2),
            "placeholder": zeroish[i],
            "placeholderPct": round(zeroish[i] / total * 100, 2),
            "distinct": distinct,
            "distinctCapped": distinct >= 3000,
            "numericRows": numeric_ok[i],
            "numericZero": numeric_zero[i],
            "min": nmin[i],
            "max": nmax[i],
            "top": [[k, v] for k, v in top],
        })

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print(f"wrote {OUT}  ({total:,} rows, {n} columns)")


if __name__ == "__main__":
    main()
