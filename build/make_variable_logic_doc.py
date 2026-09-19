# -*- coding: utf-8 -*-
"""Generate 'Varietal Engine - How Every Variable Is Used.docx'.

Every input the system takes, what it does, the arithmetic it feeds, and a
worked example using the mill's own 2026-27 figures. Also, honestly, the inputs
that are collected but currently change nothing.

    python build/make_variable_logic_doc.py
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Varietal Engine - How Every Variable Is Used.docx")

GREEN = RGBColor(0x05, 0x96, 0x69)
GREY = RGBColor(0x60, 0x60, 0x60)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def h(level, text):
    doc.add_heading(text, level=level)


def p(text, bold=False, italic=False, size=None, color=None, align=None):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        par.alignment = align
    return par


def bullet(text, b=None):
    par = doc.add_paragraph(style="List Bullet")
    if b:
        par.add_run(b).bold = True
    par.add_run(text)


def num(text, b=None):
    par = doc.add_paragraph(style="List Number")
    if b:
        par.add_run(b).bold = True
    par.add_run(text)


def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear")
    e.set(qn("w:fill"), hexc)
    tcPr.append(e)


def box(title, lines, colour="ECFDF5"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, colour)
    c.paragraphs[0].add_run(title).bold = True
    for ln in lines:
        cp = c.add_paragraph()
        r = cp.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
    doc.add_paragraph()


def note(title, lines, colour="FEF3C7"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, colour)
    c.paragraphs[0].add_run(title).bold = True
    for ln in lines:
        cp = c.add_paragraph()
        cp.add_run(ln).font.size = Pt(10)
    doc.add_paragraph()


def formula(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(10.5)
    r.font.color.rgb = GREEN
    r.bold = True


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for j, x in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(x)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        shade(c, "059669")
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            rr = cells[j].paragraphs[0].add_run(str(v))
            rr.font.size = Pt(9.5)
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    doc.add_paragraph()


# ============================== TITLE ==============================
p("HOW EVERY VARIABLE IS USED", bold=True, size=24, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
p("The complete logic of the Varietal Planning engine, with worked examples",
  size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
p("Gobind Sugar Mill, Aira - figures from the 2026-27 plot survey",
  size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

p("This document answers one question: when the system decides that a particular "
  "variety should be planted on a particular plot, what made it decide that?")
p("Every input is listed with what it controls, the arithmetic it feeds, and a "
  "worked example you can check on a calculator. The last section lists the "
  "inputs that are collected but currently change nothing - those matter just as "
  "much, because a number that looks important but does nothing is worse than a "
  "number that is missing.")

h(2, "Three kinds of variable")
table(["Kind", "Where it comes from", "Can it be argued with?"],
      [["MEASURED", "Read from the survey. Area, land type, maturity, crop stage",
        "No. It is what the surveyors found"],
       ["JUDGED", "Typed by the cane department on the Step 2 sheet",
        "Yes. It is expert opinion"],
       ["POLICY", "Set on Step 3 by the mill",
        "Yes. It is a decision, not a fact"]],
      [1.2, 3.0, 2.3])

box("THE CHAIN, IN ONE PICTURE", [
    "  MEASURED            JUDGED              POLICY",
    "  survey area   +   land suitability  +  concentration caps",
    "  land type         sucrose, yield       seed rate, multiplier",
    "  maturity          red rot, seed        ratoon assumptions",
    "        |                 |                    |",
    "        +-----------------+--------------------+",
    "                          |",
    "              1. HOW MUCH LAND IS FREE       (replant budget)",
    "              2. HOW FAR EACH VARIETY CAN GO (reachability)",
    "              3. WHICH PLOT GETS WHICH       (allocation)",
    "              4. IS IT ALLOWED               (compliance)",
    "              5. WHAT IS IT WORTH            (recovery, rupees)",
])

doc.add_page_break()

# ============================== PART A ==============================
p("PART A", bold=True, size=16, color=GREEN)
h(1, "Variables that decide HOW MUCH LAND IS FREE")

p("Nothing can be planned until we know how much land is actually available. A "
  "field carrying a plant or autumn crop becomes a ratoon of the same variety "
  "next season - its variety is already decided. Only a field finishing its "
  "ratoon is free.")

table(["Variable", "Kind", "Default", "What it does"],
      [["Ratoon : Plant ratio", "POLICY", "0.90", "Ratoon hectares standing for every 1 planted"],
       ["Ratoon II carry rate", "POLICY", "3.6%", "Share of ratoon I going to a second ratoon"],
       ["Ratoons taken", "POLICY", "2", "At 1, the carry rate is ignored entirely"],
       ["Surveyed area", "MEASURED", "56,491 ha", "The land the shares are applied to"]],
      [1.6, 0.9, 0.9, 3.0])

h(2, "The calculation")
p("Call the freshly planted area P. For every hectare planted, r hectares are "
  "standing in ratoon, and a fraction 'carry' of those go on to a second ratoon. "
  "So the total standing area is:")
formula("total = P + rP + carry x rP  =  P x (1 + r x (1 + carry))")
p("Turn it around to get the share that must be freshly planted each year:")
formula("fresh plant share = 1 / (1 + r x (1 + carry))")
formula("free to replant   = total area x (1 - fresh plant share)")

h(2, "Worked example")
box("AT THE DEFAULT SETTINGS  (r = 0.90, carry = 3.6%, 2 ratoons)", [
    "fresh share = 1 / (1 + 0.90 x 1.036)",
    "            = 1 / 1.9324",
    "            = 0.5175   ->  51.7%",
    "",
    "locked      = 56,491 x 0.5175  =  29,234 ha",
    "free        = 56,491 x 0.4825  =  27,257 ha",
])
box("AT THE SURVEY'S MEASURED VALUES  (r = 0.68, carry = 3.8%)", [
    "fresh share = 1 / (1 + 0.68 x 1.038)  =  0.586  ->  58.6%",
    "",
    "free        = 56,491 x 0.414   =  23,375 ha",
    "survey says                       23,312 ha     difference: 63 ha",
])
note("WHY THIS MATTERS MORE THAN IT LOOKS", [
    "These two settings differ by nearly 4,000 hectares of replant budget. Every "
    "downstream figure - how much seed is needed, how far a variety can expand, "
    "how many plots get instructions - scales with this one number.",
    "The screen has a 'Use measured values' button that sets r and carry from the "
    "survey in one click, and warns when the current settings sit more than 8% away "
    "from what was measured.",
])

doc.add_page_break()

# ============================== PART B ==============================
p("PART B", bold=True, size=16, color=GREEN)
h(1, "Variables that decide HOW FAR A VARIETY CAN GO")

p("A variety cannot simply be given whatever area you want. It has to be grown "
  "from seed, and seed is cane. Who supplies that cane decides what limits the "
  "growth, and the two cases are completely different.")

table(["Variable", "Kind", "Default", "What it does"],
      [["Seed rate", "POLICY", "65 qtl/ha", "Cane needed to plant one hectare"],
       ["Bud type", "POLICY", "DOUBLE", "Single bud uses about 30% less"],
       ["Multiplication factor", "POLICY", "8x", "1 ha of seed plants 8 ha next year"],
       ["Seed retention %", "JUDGED", "50%", "Share of the crop held back as seed"],
       ["Seed available", "JUDGED", "-", "Quintals the MILL holds of this variety"],
       ["Seed purchase ceiling", "POLICY", "35 ha", "Most that can be bought from the breeder"],
       ["Current area", "MEASURED", "-", "Decides which of the two cases applies"]],
      [1.6, 0.9, 0.9, 3.0])

h(2, "The two cases")
box("WHO SUPPLIES THE SEED DECIDES THE CEILING", [
    "FARMERS ALREADY GROW IT   they cut seed from their own standing crop, for",
    "  (current area > 0)      replanting and for new ground alike. The mill's",
    "                          nursery does not limit them. Their own",
    "                          multiplication rate does.",
    "",
    "NOT OUT THERE YET         every hectare must come from mill seed, so the",
    "  (current area = 0)      nursery stock and purchase ceiling bind hard.",
])

h(2, "The calculation")
p("Established variety:", bold=True)
formula("expansion ratio = 1 + (multiplier x retention - 1) x 0.35")
formula("max reachable   = current area x expansion ratio")
p("New variety:", bold=True)
formula("effective seed rate = double bud ? 65 : 65 x 0.7 = 45.5 qtl/ha")
formula("max reachable = min( seed held / effective rate , purchase ceiling )")

h(2, "Worked example - an established variety")
box("COS 13231, standing on 2,062 ha, multiplier 8", [
    "retention  25%:  ratio = 1 + (8 x 0.25 - 1) x 0.35 = 1.35   ->  2,784 ha",
    "retention  50%:  ratio = 1 + (8 x 0.50 - 1) x 0.35 = 2.05   ->  4,227 ha",
    "retention 100%:  ratio = 1 + (8 x 1.00 - 1) x 0.35 = 3.45   ->  7,114 ha",
])
p("This is the seed-versus-sugar trade-off priced. Holding back more cane as seed "
  "reaches further, but that cane is not crushed, so this year's sugar is lower.")

h(2, "Worked example - a new variety")
box("A VARIETY NOBODY GROWS YET, 3,000 qtl of mill seed, double bud", [
    "plantable from stock = 3,000 / 65      = 46 ha",
    "purchase ceiling                       = 35 ha",
    "max reachable        = min(46, 35)     = 35 ha",
    "",
    "next year  35 x 8 = 280 ha        year after  280 x 8 = 2,240 ha",
])
note("SEED IS ONLY REQUIRED FOR GROUND A VARIETY DOES NOT ALREADY HOLD", [
    "When a farmer replants the same variety on the same land he cuts seed from his "
    "own crop. He does not buy it.",
    "So the mill's seed requirement counts only NEW area. Charging every variety for "
    "all of its replanting demanded 1,515,280 quintals across the command area - "
    "about 26 times what the mill holds - and put a red warning on every variety on "
    "the screen. Counting growth alone brings it to a figure you can act on.",
])

doc.add_page_break()

# ============================== PART C ==============================
p("PART C", bold=True, size=16, color=GREEN)
h(1, "Variables that decide WHICH PLOT GETS WHICH VARIETY")

p("This is the part that turns a plan into an instruction. Every plot finishing "
  "its ratoon - 62,717 of them, 23,161 hectares - is offered to the varieties "
  "that are allowed to take it.")

table(["Variable", "Kind", "What it does"],
      [["Land suitability", "JUDGED", "UPLAND / LOWLAND / BOTH. A variety may only take matching land"],
       ["Plot land type", "MEASURED*", "What the plot actually is. *Inferred on ratoon plots - see below"],
       ["Maturity", "MEASURED", "EARLY / GENERAL / REJECTED, read from the survey's CROPCATEGORY"],
       ["Plot area", "MEASURED", "Decides whether the plot can be split between varieties"],
       ["Village cap", "POLICY", "No village may exceed this share on one variety"],
       ["Strategy", "JUDGED", "An EXIT variety is never planted, and its plots are freed first"]],
      [1.5, 1.0, 4.0])

h(2, "Rule 1 - a variety may only take land it suits")
box("ELIGIBILITY", [
    "variety says BOTH      -> may take any plot",
    "variety says UPLAND    -> only upland plots",
    "variety says LOWLAND   -> only lowland plots",
    "variety says UNKNOWN   -> NOT PLACED AT ALL",
    "plot land type UNKNOWN -> only BOTH varieties may take it",
])
p("The last two lines are deliberate. A variety nobody has classified is not "
  "quietly assumed to be upland, and neither is a plot. Guessing is how 23,161 "
  "hectares came to be treated as upland when roughly 40% of it is lowland.")

h(2, "The land type problem, and how it is solved")
note("THE ERP RECORDS NO LAND TYPE ON RATOON ROWS", [
    "All 92,319 RATOON rows come through blank and default to UPLAND. Since every "
    "free plot is a ratoon plot, the allocator was choosing varieties for 23,161 ha "
    "it believed was entirely upland.",
    "Proof that it is a recording gap and not reality: RATOON II rows, which the ERP "
    "does fill in, are 59% lowland. Fields do not become upland when they ratoon.",
])
p("Each ratoon plot's land type is inferred from land the survey DID record:")
table(["Source", "How", "Share of free area"],
      [["MEASURED", "Ratoon II - the ERP recorded it, used as-is", "3.6%"],
       ["GROWER", "This grower's own surveyed plots give his lowland share", "95.2%"],
       ["VILLAGE", "No measured plot for this grower, village share used", "1.1%"],
       ["UNKNOWN", "Neither available - left alone, never guessed", "0%"]],
      [1.2, 3.8, 1.5])

p("Within a grower the assignment is proportional, not plot by plot:")
formula("lowland budget = total free area of his plots x his measured lowland share")
p("His plots are taken largest first, and a plot is marked lowland while most of "
  "it still fits inside the remaining budget.")

box("WORKED EXAMPLE - one grower", [
    "His surveyed plant crop:  1.2 ha,  of which 0.48 ha lowland  ->  share 40%",
    "His free ratoon plots:    0.6 ha,  0.4 ha,  0.3 ha           ->  total 1.3 ha",
    "",
    "lowland budget = 1.3 x 0.40 = 0.52 ha",
    "",
    "  plot 0.6 ha : budget 0.52 >= 0.30 (half of it)  -> LOWLAND, budget 0.52-0.6 = -0.08",
    "  plot 0.4 ha : budget -0.08 < 0.20               -> UPLAND",
    "  plot 0.3 ha : budget -0.08 < 0.15               -> UPLAND",
    "",
    "result: 0.6 of 1.3 ha lowland = 46%, against his measured 40%",
])
p("Across all 62,717 plots this produces 41.4% lowland, against the 39.5% measured "
  "on surveyed rows. The two are computed from different plots, so the agreement "
  "is a check on the method rather than a restatement of it.")

h(2, "Rule 2 - a big plot may carry more than one variety")
formula("blocks a plot can carry = floor( plot area / 0.25 ha )")
p("A block below a quarter hectare cannot be harvested separately, so a plot "
  "smaller than 0.5 ha stays on one variety. At Gobind the median plot is 0.23 ha "
  "and only 15% reach half a hectare - but those 15% hold 40% of the area.")
box("WORKED EXAMPLE - splitting", [
    "plot 0.23 ha  ->  floor(0.23 / 0.25) = 0  ->  1 block   (one variety)",
    "plot 0.60 ha  ->  floor(0.60 / 0.25) = 2  ->  2 blocks  (may be split)",
    "plot 1.40 ha  ->  floor(1.40 / 0.25) = 5  ->  capped at 2 in practice",
])
p("Each block ratoons as its own variety, so the variety lock is finer than the "
  "plot. In the last run 8,158 plots were split between two varieties.")

h(2, "Rule 3 - a grower's plots are filled together, for staggered harvest")
p("92.6% of the command area is EARLY maturity, so the whole crop peaks at once. "
  "The allocator therefore prefers to give a grower's next plot a maturity he has "
  "not been given yet. Because the median plot is too small to split, this is how "
  "most staggering is achieved - across his plots rather than within one. 69% of "
  "growers have two or more plots.")

h(2, "Rule 4 - no village may be taken over by one variety")
formula("village headroom = village free area x village cap % - already given")
p("Enforced while assigning, not checked afterwards. A variety with no headroom "
  "left in a village is simply not offered its plots.")
box("WORKED EXAMPLE - village cap at 60%", [
    "village has 100 ha free, CO 0118 already given 55 ha",
    "",
    "cap      = 100 x 0.60 = 60 ha",
    "headroom = 60 - 55    =  5 ha",
    "",
    "-> CO 0118 may take at most 5 more ha here; the rest goes to another variety",
])

h(2, "Rule 5 - plots on an exit variety are freed first")
p("Within each grower, plots finishing on a variety the plan wants gone are "
  "offered before the others. Those are the plots the plan most wants to move. "
  "Larger plots come next, since they can carry a split.")

doc.add_page_break()

# ============================== PART D ==============================
p("PART D", bold=True, size=16, color=GREEN)
h(1, "Variables that decide WHETHER THE PLAN IS ALLOWED")

p("Six checks run on every plan. Each is a rule the mill set on Step 3, tested "
  "against what the plan produces.")

table(["Check", "Variable", "Test"],
      [["Concentration cap", "Max variety concentration (40%)",
        "No variety above 40% of the command area"],
       ["Village cap", "Village concentration cap (60%)",
        "No village above 60% on one variety"],
       ["Lowland floor", "Lowland coverage floor (28%)",
        "Lowland-capable varieties at or above 28%"],
       ["Red rot", "Red rot trigger (2%)",
        "Area on susceptible varieties at or below 2%"],
       ["Seed sufficiency", "Seed available, seed rate",
        "No variety short of seed for its growth"],
       ["Replant fit", "Ratoon assumptions",
        "Fresh planting fits the land coming out of ratoon"]],
      [1.4, 2.1, 3.0])

h(2, "Worked example - the concentration cap")
formula("cap in hectares = 40% x 56,491 = 22,596 ha")
box("CO 0118", [
    "current area   18,029 ha  ->  31.9% of the command area",
    "cap            22,596 ha",
    "headroom        4,567 ha   ->  PASSES, but not by a wide margin",
])

h(2, "Worked example - the red rot check")
p("Step 2's sheet carries a red rot reaction per variety. Any variety marked S is "
  "a standing outbreak risk.")
formula("susceptible % = area on S varieties / total planned area")
box("IF CO 0118 WERE MARKED S", [
    "susceptible area = 18,029 ha of 29,351 ha planned  =  61.4%",
    "trigger          = 2%",
    "-> FAILS. The plan rests 61% of its area on a susceptible variety.",
    "",
    "With no red rot recorded at all, the check reports 'Not yet known'",
    "rather than passing quietly.",
])

h(2, "Worked example - reachability, which catches an impossible target")
box("COS 13231 asked for 9,000 ha in Year 1", [
    "current area        2,062 ha",
    "retention              50%",
    "max reachable       4,227 ha",
    "shortfall           4,773 ha   ->  FAILS",
])
p("The check judges what was ASKED FOR, not what survived the caps. A target of "
  "60,000 ha would otherwise be trimmed to the 22,596 ha cap first, pass the "
  "reachability test, and tell the planner nothing at all.")

doc.add_page_break()

# ============================== PART E ==============================
p("PART E", bold=True, size=16, color=GREEN)
h(1, "Variables that decide WHAT THE PLAN IS WORTH")

table(["Variable", "Kind", "Default", "What it does"],
      [["Juice sucrose %", "JUDGED", "-", "Sugar in the cane, per variety"],
       ["Juice to recovery factor", "POLICY", "0.635", "Converts sucrose to recovery"],
       ["Season crush", "POLICY", "13.5 lakh MT", "Cane crushed, for costing a recovery change"],
       ["Sugar price", "POLICY", "Rs 38/kg", "For the same purpose"],
       ["Cane yield t/ha", "JUDGED", "-", "Feeds seed generated; see the gap below"]],
      [1.7, 0.9, 1.2, 2.7])

h(2, "The calculation")
formula("blended sucrose = sum(area x sucrose) / total area")
formula("recovery        = blended sucrose x 0.635")
formula("extra sugar MT  = season crush x recovery gain %")
formula("value in crores = extra sugar MT x 1000 x price / 1,00,00,000")

h(2, "Worked example")
box("EXPANDING THE HIGH-SUGAR VARIETIES", [
    "base blended sucrose   18.24%   ->  recovery 11.58%",
    "year 3 blended sucrose 18.53%   ->  recovery 11.77%",
    "recovery gain                        0.19%",
    "",
    "extra sugar = 13,50,000 x 0.0019          =  2,565 MT",
    "value       = 2,565 x 1000 x 38 / 1e7     =  Rs 9.7 crores",
])
note("THE SAME FACTOR IS USED FOR THE BASE YEAR AND EVERY PROJECTED YEAR", [
    "It did not used to be. The base year used 0.635 and projections used 0.638, "
    "and projections were additionally given 0.08% recovery per year automatically - "
    "Year 3 opened 0.24% ahead before any varietal change was counted.",
    "A plan that changed nothing showed a recovery gain and a crores-level benefit. "
    "That has been removed. Recovery now moves only when the varietal mix moves: "
    "leave every variety on HOLD and the gain reads +0.00%.",
])
note("THE BIGGEST REMAINING GAP: THERE IS NO TONNAGE", [
    "Recovery is a percentage - sugar per tonne of cane. The value calculation uses "
    "a flat 13.5 lakh MT rather than area multiplied by each variety's yield.",
    "So the system can say a plan improves recovery, but not that it also changes "
    "how much cane there is to crush. A variety with 18.5% sucrose and poor yield "
    "can lose to one with 17% sucrose and heavy cane, and the system cannot "
    "currently see that. Cane yield per variety is the input that would close it.",
])

doc.add_page_break()

# ============================== FULL EXAMPLE ==============================
h(1, "One plot, end to end")

p("A single plot, followed from the survey to the instruction a supervisor "
  "receives.")

box("WHAT THE SURVEY SAYS", [
    "village        RAMPUR",
    "grower         12345",
    "area           0.60 ha",
    "crop type      RATOON          -> finishing, so FREE to replant",
    "variety        CO 0238         -> the plan has this variety on EXIT",
    "land type      (blank)         -> the ERP records none on ratoon rows",
])
box("STEP 1 - WHAT IS THIS PLOT?", [
    "free to replant   yes, it is finishing ratoon",
    "land type         unknown, so inferred:",
    "                  grower 12345's surveyed plots are 40% lowland",
    "                  his free plots total 1.3 ha -> lowland budget 0.52 ha",
    "                  this is his largest plot at 0.60 ha -> LOWLAND (source: GROWER)",
])
box("STEP 2 - WHO IS ELIGIBLE?", [
    "CO 0118      BOTH      -> eligible",
    "COLK 94184   LOWLAND   -> eligible",
    "CO 15023     UPLAND    -> NOT eligible, this plot is lowland",
    "CO 0238      EXIT      -> never planted",
    "COS 8272     UNKNOWN   -> not placed, nobody has classified it",
])
box("STEP 3 - WHICH OF THEM, AND HOW MANY?", [
    "plot is 0.60 ha  ->  floor(0.60 / 0.25) = 2 blocks, so it may be SPLIT",
    "",
    "this plot is on an EXIT variety, so it is offered FIRST among this grower's",
    "plots - the plan most wants CO 0238 gone",
    "",
    "block 1: COLK 94184 needs the most area and has village headroom  -> 0.30 ha",
    "block 2: CO 0118 is EARLY, COLK 94184 is EARLY too, so no maturity",
    "         gain here; CO 0118 taken on remaining need                -> 0.30 ha",
])
box("THE INSTRUCTION", [
    "RAMPUR / grower 12345 / plot 0.60 ha / LOWLAND (inferred from grower)",
    "",
    "   plant COLK 94184 on 0.30 ha",
    "   plant CO 0118    on 0.30 ha",
    "   replaces CO 0238, which the plan is exiting     [PRIORITY]",
])

doc.add_page_break()

# ============================== NOT USED ==============================
h(1, "Collected, but currently changes nothing")

p("These are real inputs on the Step 2 sheet or the screens. They are stored and "
  "displayed, but no calculation reads them. They are listed here so nobody "
  "assumes the plan responds to them.")

table(["Variable", "Where it is collected", "Status"],
      [["Crop duration (12 / 18 month)", "Step 2 sheet",
        "Read into the record, then used nowhere at all"],
       ["Planting season", "Step 2 sheet and cards",
        "Shown on the card; no calculation uses it"],
       ["Avg cane weight", "Step 2 sheet and cards", "Display only"],
       ["Animal damage risk", "Step 2 sheet and cards", "Display only"],
       ["Farmer acceptance", "Step 2 sheet and cards", "Display only"],
       ["Measured lowland %", "Derived from the survey",
        "Carried on the record but never read; the cross-check is on the Excel sheet only"]],
      [1.9, 2.0, 2.6])

note("WHY CROP DURATION MATTERS MOST OF THESE", [
    "An 18-month autumn crop occupies its field through two seasons; a 12-month "
    "spring crop through one. That difference belongs in the replant budget, which "
    "currently applies one ratio to every variety alike.",
    "It is collected on the sheet and would change the answer, but nothing reads it "
    "yet. Until it does, a mill planting a lot of autumn cane will see a replant "
    "budget that is too generous.",
])

h(2, "Two constants nobody has explained")
table(["Constant", "Where", "Effect"],
      [["0.35", "Expansion ratio in Part B",
        "Turns 8x multiplication into a 2.05x area ratio. Sets every reachability figure on screen"],
       ["0.15", "Seed plot fraction",
        "Share of a variety's area treated as seed plot when computing next year's stock"]],
      [1.0, 2.2, 3.3])
p("Both change results and neither has a stated source. They are on the list of "
  "questions for the plant team head.")

doc.add_page_break()

# ============================== QUICK REF ==============================
h(1, "Quick reference - every formula in one place")

box("LAND", [
    "fresh plant share = 1 / (1 + r x (1 + carry))",
    "free to replant   = total area x (1 - fresh plant share)",
])
box("SEED", [
    "effective seed rate = double bud ? rate : rate x 0.7",
    "seed required qtl   = new area x effective seed rate",
    "plantable ha        = floor(seed qtl / effective seed rate)",
])
box("REACH", [
    "expansion ratio = 1 + (multiplier x retention - 1) x 0.35",
    "established     = current area x expansion ratio",
    "new variety     = min(seed held / rate, purchase ceiling)",
])
box("ALLOCATION", [
    "blocks per plot  = floor(plot area / 0.25 ha)",
    "village headroom = village free x cap% - already given",
    "lowland budget   = grower free area x his measured lowland share",
])
box("VALUE", [
    "blended sucrose = sum(area x sucrose) / total area",
    "recovery        = blended sucrose x 0.635",
    "extra sugar MT  = season crush x recovery gain %",
    "value crores    = extra sugar MT x 1000 x price / 1e7",
])
box("CAPS", [
    "variety cap ha = max concentration % x command area",
    "susceptible %  = area on red-rot-S varieties / planned area",
])

doc.add_paragraph()
p("Generated from the 2026-27 plot survey. Every figure can be checked against "
  "that file.", size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print(f"wrote {OUT}")
print(f"  {os.path.getsize(OUT)/1024:.0f} KB")
