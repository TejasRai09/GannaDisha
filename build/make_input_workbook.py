# -*- coding: utf-8 -*-
"""Build the single input workbook the cane department fills in.

Everything the planning system needs a human to decide, in one file:

    Start Here          what to do, in what order, and what each tab is for
    A. Varieties        what each variety IS - 86 rows, already listed
    B. Agronomic Rules  the mill's own settings, with the survey's measured
                        value beside each so a choice can be made on evidence
    C. Strategy         where each variety should go, and how much cane to
                        hold back as seed to get it there
    Reference           accepted values and a short glossary

Measured columns are pre-filled and locked by convention. Judgement columns are
blank with dropdowns, so a typo cannot enter the engine.

    python build/make_input_workbook.py
"""

import datetime
import os
import subprocess
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CACHE = os.path.join(HERE, "cache", "plots_2627.parquet")
OUT_DIR = os.path.join(ROOT, "plan_inputs")
OUT = os.path.join(OUT_DIR, "Varietal Plan - Input Workbook.xlsx")

CORE_COVERAGE = 0.95

INK = "1F2937"
HDR_LOCKED = "334155"
HDR_FILL = "0F766E"
BAND_LOCKED = "F1F5F9"
BAND_CORE = "FEF9C3"
TITLE_BG = "0F766E"
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LISTS = {
    "Stage": ["TRIAL", "MULTIPLYING", "COMMERCIAL", "DECLINING", "RETIRED", "REVIEW"],
    "Strategy": ["INTRODUCE-NEW", "EXPAND", "HOLD", "REDUCE", "EXIT"],
    "Land Suitability": ["UPLAND", "LOWLAND", "BOTH"],
    "Planting Season": ["SPRING", "AUTUMN", "BOTH"],
    "Crop Duration": ["12-MONTH", "18-MONTH"],
    "Red Rot Resistance": ["R", "MR", "S"],
    "Animal Damage Risk": ["LOW", "MEDIUM", "HIGH"],
    "Seed Retention": ["100% (all cane kept as seed)", "50% (half)", "25% (mostly crushed)"],
}


# --------------------------------------------------------------- helpers
def title_block(ws, ncol, title, subtitle, note):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TITLE_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(row=2, column=2 - 1, value=subtitle)
    c.font = Font(size=10, italic=True, color="475569")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.cell(row=3, column=1, value=note).font = Font(size=9, color="94A3B8")


def header_row(ws, row, cols, locked_upto):
    for i, (label, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HDR_LOCKED if i <= locked_upto else HDR_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 34


def dropdown(ws, col_letter, first, last, options, label):
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showDropDown=False,
        errorTitle="Not an accepted value",
        error="Choose one of: " + ", ".join(options),
        promptTitle=label,
        prompt="Pick one of: " + ", ".join(options),
    )
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first}:{col_letter}{last}")


def guard(ws, col_letter, first, last, lo, hi, label):
    dv = DataValidation(
        type="decimal", operator="between", formula1=lo, formula2=hi, allow_blank=True,
        errorTitle="Out of range",
        error=f"{label} should be between {lo} and {hi}. Leave blank if not known.",
    )
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first}:{col_letter}{last}")


def load_varieties():
    if not os.path.exists(CACHE):
        subprocess.run([sys.executable, os.path.join(HERE, "extract_2627.py")], check=True)
    df = pd.read_parquet(CACHE)
    total = df.area_ha.sum()
    g = (df.groupby("variety")
           .agg(area_ha=("area_ha", "sum"), records=("variety", "size"))
           .reset_index())
    rec = df[df.crop_type != "RATOON"]
    low = rec.assign(_l=rec.land_type.eq("LOWLAND")).groupby("variety")._l.mean()
    mat = df.groupby("variety").crop_category.agg(
        lambda x: x.mode().iat[0] if len(x.mode()) else "UNKNOWN")
    g["lowland_pct"] = (g.variety.map(low) * 100).fillna(0)
    g["maturity"] = g.variety.map(mat).fillna("UNKNOWN")
    g = g[g.area_ha > 0].sort_values("area_ha", ascending=False).reset_index(drop=True)
    g["share"] = g.area_ha / total * 100
    g["cum"] = g.share.cumsum()
    g["priority"] = ["FILL FIRST" if c <= CORE_COVERAGE * 100 else "optional" for c in g.cum]
    return g, total, df


# --------------------------------------------------------------- sheets
def sheet_start(wb, nvar, core, total, stamp):
    ws = wb.create_sheet("Start Here")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 96

    ws.merge_cells("A1:C1")
    c = ws.cell(row=1, column=1, value="VARIETAL PLAN - INPUT WORKBOOK")
    c.font = Font(bold=True, size=16, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TITLE_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:C2")
    c = ws.cell(row=2, column=1, value="Gobind Sugar Mill, Aira  -  everything the planning system needs a person to decide")
    c.font = Font(size=10, italic=True, color="475569")
    c.alignment = Alignment(horizontal="center")

    r = [4]

    def line(b, text="", bold_b=True, size=10, gap=False):
        if gap:
            r[0] += 1
        ws.cell(row=r[0], column=2, value=b).font = Font(bold=bold_b, size=size, color=INK)
        cell = ws.cell(row=r[0], column=3, value=text)
        cell.font = Font(size=size, color=INK)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r[0] += 1

    line("What this file is for", "", size=12)
    line("", "The survey tells us where every field is and what is growing on it. It cannot tell us what a variety is LIKE, "
             "what rules the mill wants to plan by, or where each variety should go next. Those three things are what this "
             "workbook collects. Fill it in, send it back, and the system produces the three-year plan.")
    line("What is already done for you", "", size=12, gap=True)
    line("", f"All {nvar} varieties the {stamp} survey found are already listed, biggest area first. You do not have to type "
             "a single variety name. The grey columns are measured facts from the survey - please do not change them.")
    line("The three tabs, in order", "", size=12, gap=True)
    line("A. Varieties", "WHAT each variety is. Its land, its season, its sugar, its disease reaction, and how much seed we hold. "
                         "This is the longest tab but only the FILL FIRST rows really matter.")
    line("B. Agronomic Rules", "HOW the mill wants to plan. Seed rate, multiplication, ratoon assumptions, and the safety caps. "
                               "Each one shows what the survey actually measured beside it, so you can decide on evidence rather than memory.")
    line("C. Strategy", "WHERE each variety should go - expand, hold, reduce or exit - and how much cane to hold back as seed to get it there.")
    line("Two rules that matter", "", size=12, gap=True)
    line("1.  Blank is not zero", "If you do not know a number, LEAVE IT BLANK. Do not put 0. Blank means 'nobody has told us yet' "
                                  "and the system will ask for it. A 0 means 'we measured it and it is zero', which is a different thing "
                                  "and will quietly spoil the plan.")
    line("2.  Seed Available is the exception", "There, 0 is a real and useful answer - it means the mill holds no seed of that variety.")
    line("Where to start", "", size=12, gap=True)
    line("", f"Tab A, the {core} rows marked FILL FIRST. They cover 95% of the command area. "
             f"The other {nvar - core} varieties average about 44 ha each and can wait.")
    line("A useful cross-check", "", size=12, gap=True)
    line("", "On tab A, the column 'Lowland % (measured)' sits right beside 'Land Suitability'. One is what the survey found, the "
             "other is what you say. If a variety reads 37% lowland and it is marked UPLAND, one of the two is wrong - worth a "
             "look before the sheet comes back.")
    line("If anything is unclear", "", size=12, gap=True)
    line("", "The Reference tab lists every accepted value and explains the terms. Anything still unclear, leave blank and add a "
             "note - a blank the system asks about is far better than a guess it believes.")

    ws.cell(row=r[0] + 1, column=2, value=f"Survey: 2026-27 plot-wise report  |  {total:,.1f} ha  |  generated {stamp}").font = \
        Font(size=9, color="94A3B8")
    return ws


def sheet_varieties(wb, g):
    cols = [
        ("#", 5), ("Variety", 22), ("Area (ha)", 12), ("% of Area", 10),
        ("Lowland % (measured)", 19), ("Maturity (measured)", 18), ("Priority", 12),
        ("Stage", 15), ("Land Suitability", 17), ("Planting Season", 16),
        ("Crop Duration", 14), ("Juice Sucrose %", 15), ("Cane Yield (t/ha)", 16),
        ("Avg Cane Weight (g)", 18), ("Red Rot Resistance", 18),
        ("Animal Damage Risk", 18), ("Farmer Acceptance (1-5)", 21),
        ("Seed Available (qtl)", 18), ("Notes", 42),
    ]
    LOCKED = 7
    ws = wb.create_sheet("A. Varieties")
    title_block(ws, len(cols), "TAB A - WHAT EACH VARIETY IS",
                "Grey columns come from the survey. Fill the teal ones. Start with the FILL FIRST rows.",
                "Leave a cell blank if you do not know it. Do not put 0 - see Start Here.")
    HDR = 5
    header_row(ws, HDR, cols, LOCKED)
    first = HDR + 1

    for i, row in enumerate(g.itertuples()):
        r = first + i
        vals = [i + 1, row.variety, round(float(row.area_ha), 1), round(float(row.share), 2),
                round(float(row.lowland_pct), 1), row.maturity, row.priority]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = Font(size=10, color=INK, bold=(j == 2))
            fill = BAND_CORE if (j == 7 and row.priority == "FILL FIRST") else BAND_LOCKED
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(horizontal="right" if j in (1, 3, 4, 5) else "left")
        for j in range(LOCKED + 1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.border = BORDER
            c.font = Font(size=10, color=INK)

    last = first + len(g) - 1
    for name, letter in [("Stage", "H"), ("Land Suitability", "I"), ("Planting Season", "J"),
                         ("Crop Duration", "K"), ("Red Rot Resistance", "O"), ("Animal Damage Risk", "P")]:
        dropdown(ws, letter, first, last, LISTS[name], name)
    for letter, lo, hi, label in [("L", 10, 24, "Juice sucrose %"), ("M", 20, 160, "Cane yield t/ha"),
                                  ("N", 200, 2000, "Cane weight g"), ("Q", 1, 5, "Farmer acceptance"),
                                  ("R", 0, 10_000_000, "Seed available")]:
        guard(ws, letter, first, last, lo, hi, label)

    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{last}"
    return ws


def sheet_rules(wb, measured):
    cols = [("Setting", 30), ("What it controls", 54), ("Current default", 15),
            ("Survey measured", 16), ("YOUR VALUE", 14), ("Notes", 40)]
    ws = wb.create_sheet("B. Agronomic Rules")
    title_block(ws, len(cols), "TAB B - HOW THE MILL WANTS TO PLAN",
                "Each setting shows what the survey actually measured, so the choice can rest on evidence.",
                "Leave YOUR VALUE blank to keep the current default.")
    HDR = 5
    header_row(ws, HDR, cols, 4)

    groups = [
        ("SEED AND MULTIPLICATION", [
            ("Seed rate (qtl/ha)", "Cane needed to plant one hectare. 65 qtl = 6.5 tonnes.", "65", "-", ""),
            ("Bud type", "SINGLE BUD uses about 30% less cane than DOUBLE BUD.", "DOUBLE BUD", "-", ""),
            ("Multiplication factor", "1 ha of seed plants how many ha next year. It is really the ratio of nursery to farmer planting density.", "8", "-", "Confirm the two spacings"),
            ("Seed purchase ceiling (ha)", "Most that can be bought from the breeding institute in one year.", "35", "-", ""),
            ("Test plot size (ha)", "Size of a first trial of a new variety.", "5", "-", ""),
        ]),
        ("CROP CYCLE - decides how much land is free each year", [
            ("Ratoons taken", "How many ratoons before ploughing out. At 1, the carry rate below is ignored.", "2", "effectively 1", "Only 3.8% reach a second ratoon"),
            ("Ratoon : Plant ratio", "Ratoon hectares standing for every 1 hectare planted. This one number moves the replant budget by ~4,000 ha.", "0.90", f"{measured['ratio']:.2f}", "0.90 is the guiding document's target; 0.68 is what happens"),
            ("Ratoon II carry rate (%)", "Share of ratoon I fields kept for a second ratoon.", "3.6", f"{measured['carry']:.1f}", "Default is close to measured"),
        ]),
        ("SAFETY CAPS - manager owned", [
            ("Max variety concentration (%)", "No single variety above this share of the command area.", "40", f"{measured['top_share']:.1f} now", f"{measured['top_name']} is the largest"),
            ("Village concentration cap (%)", "No single village above this share on one variety. This is where an outbreak starts.", "60", "-", ""),
            ("Lowland coverage floor (%)", "Minimum share on lowland-capable varieties.", "28", f"{measured['lowland']:.1f}", "Measured is above the floor - no breach"),
            ("Red rot trigger (%)", "Most area allowed on red-rot susceptible varieties before the plan is flagged.", "2", "-", "Needs the Red Rot column on tab A"),
        ]),
        ("MILL OUTPUT - decides what a plan is worth", [
            ("Juice to recovery factor", "Recovery as a share of juice sucrose. 18% sucrose x 0.635 = 11.43% recovery.", "0.635", "-", "Mill's own figure"),
            ("Season crush (MT cane)", "Cane crushed in a season, used to price a recovery change.", "13,50,000", "-", ""),
            ("Sugar price (Rs/kg)", "Ex-mill price, used for the same purpose.", "38", "-", ""),
        ]),
        ("SCOPE", [
            ("Command area (ha)", "Total land planned for.", "57,000", f"{measured['total']:,.0f}", "Survey is within 1%"),
            ("Planning horizon (years)", "How many years the plan covers. Can be 1 to 5.", "3", "-", ""),
            ("Base year", "The season the survey describes.", "2026-27", "-", ""),
        ]),
    ]

    r = HDR + 1
    first_input = r
    for gname, rows in groups:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        c = ws.cell(row=r, column=1, value=gname)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="475569")
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 20
        r += 1
        for name, what, default, meas, note in rows:
            for j, v in enumerate([name, what, default, meas, None, note], start=1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = BORDER
                c.font = Font(size=10, color=INK, bold=(j == 1))
                c.alignment = Alignment(wrap_text=(j in (2, 6)), vertical="top",
                                        horizontal="center" if j in (3, 4, 5) else "left")
                if j <= 4:
                    c.fill = PatternFill("solid", fgColor=BAND_LOCKED)
            ws.row_dimensions[r].height = 30
            r += 1

    ws.freeze_panes = ws.cell(row=first_input, column=1)
    return ws


def sheet_strategy(wb, g):
    cols = [("#", 5), ("Variety", 22), ("Area (ha)", 12), ("% of Area", 10),
            ("Maturity", 14), ("Priority", 12),
            ("Strategy", 18), ("Seed Retention", 28), ("Target Y1 (ha)", 14),
            ("Target Y2 (ha)", 14), ("Target Y3 (ha)", 14), ("Notes", 40)]
    LOCKED = 6
    ws = wb.create_sheet("C. Strategy")
    title_block(ws, len(cols), "TAB C - WHERE EACH VARIETY SHOULD GO",
                "Strategy is the direction. Seed retention is how much cane is held back to get there.",
                "Targets are optional - fill them only if a specific hectarage is wanted.")
    HDR = 5
    header_row(ws, HDR, cols, LOCKED)
    first = HDR + 1

    for i, row in enumerate(g.itertuples()):
        r = first + i
        vals = [i + 1, row.variety, round(float(row.area_ha), 1), round(float(row.share), 2),
                row.maturity, row.priority]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = Font(size=10, color=INK, bold=(j == 2))
            fill = BAND_CORE if (j == 6 and row.priority == "FILL FIRST") else BAND_LOCKED
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(horizontal="right" if j in (1, 3, 4) else "left")
        for j in range(LOCKED + 1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.border = BORDER
            c.font = Font(size=10, color=INK)

    last = first + len(g) - 1
    dropdown(ws, "G", first, last, LISTS["Strategy"], "Strategy")
    dropdown(ws, "H", first, last, LISTS["Seed Retention"], "Seed Retention")
    for letter in ("I", "J", "K"):
        guard(ws, letter, first, last, 0, 60000, "Target hectares")

    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{last}"
    return ws


def sheet_reference(wb):
    ws = wb.create_sheet("Reference")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96
    title_block(ws, 2, "REFERENCE", "Accepted values and what the terms mean", "")

    r = [5]

    def sect(t):
        c = ws.cell(row=r[0], column=1, value=t)
        c.font = Font(bold=True, size=11, color="0F766E")
        r[0] += 1

    def line(a, b):
        ws.cell(row=r[0], column=1, value=a).font = Font(bold=True, size=10)
        c = ws.cell(row=r[0], column=2, value=b)
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0] += 1

    sect("ACCEPTED VALUES")
    for k, v in LISTS.items():
        line(k, " / ".join(v))
    r[0] += 1

    sect("WHAT THE COLUMNS ON TAB A MEAN")
    for a, b in [
        ("Stage", "Where the variety is in its life. REVIEW if you are not sure yet - a small area may be a trial on its way up, not a failure."),
        ("Land Suitability", "Which land it belongs on. BOTH if it grows anywhere. Compare against the measured lowland % beside it."),
        ("Planting Season", "SPRING is Feb-Mar, AUTUMN is Oct-Nov."),
        ("Crop Duration", "12-MONTH or 18-MONTH. An 18-month autumn crop holds the field through two seasons, which changes how much land is free next year."),
        ("Juice Sucrose %", "Pol in cane, roughly 14 to 20. Leave blank if not known."),
        ("Cane Yield (t/ha)", "Tonnes per hectare, roughly 50 to 110. This decides how much sugar a plan is actually worth, and it is the figure we are most short of."),
        ("Avg Cane Weight (g)", "Single cane weight. This is what the farmer notices."),
        ("Red Rot Resistance", "R resistant, MR moderately resistant, S susceptible. An S variety should not be expanded."),
        ("Animal Damage Risk", "Soft sweet canes get eaten. LOW, MEDIUM or HIGH."),
        ("Farmer Acceptance", "1 means farmers refuse it, 5 means they ask for it. A plan farmers reject is not a plan."),
        ("Seed Available (qtl)", "Quintals of seed the MILL holds. Only matters for a variety farmers do not already grow - once it is out there they cut seed from their own crop."),
    ]:
        line(a, b)
    r[0] += 1

    sect("WHAT THE COLUMNS ON TAB C MEAN")
    for a, b in [
        ("Strategy", "EXPAND grow it, HOLD keep it steady, REDUCE bring it down, EXIT stop it, INTRODUCE-NEW start it from a nursery."),
        ("Seed Retention", "Cane kept back as seed cannot be crushed. 100% reaches furthest but gives no sugar from that variety this year. 50% is the usual balance."),
        ("Target Y1/Y2/Y3", "Only fill these if a specific hectarage is wanted. Left blank, the system works out how far seed can take each variety on its own."),
    ]:
        line(a, b)
    r[0] += 1

    sect("GLOSSARY")
    for a, b in [
        ("Plant crop", "A field grown from setts you planted."),
        ("Ratoon", "Regrowth from the stubble of a harvested crop. Same roots, same variety, no replanting - which is why a ratoon field's variety cannot be changed."),
        ("Sett", "A piece of cane with a bud on it, used instead of seed."),
        ("Locked area", "Land whose variety cannot be changed next year because a crop is standing on it."),
        ("Free to replant", "Land finishing its cycle - the only land a new decision can touch. At Gobind this is about 23,300 ha of 56,491."),
        ("Upland / Lowland", "Well-drained land versus land that waterlogs and needs tolerant varieties."),
        ("Quintal (qtl)", "100 kilograms."),
        ("Recovery", "Sugar actually extracted, as a share of cane crushed."),
        ("Maturity", "EARLY / GENERAL / REJECTED - the UP variety class, read from the survey."),
    ]:
        line(a, b)
    return ws


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    g, total, df = load_varieties()
    stamp = datetime.datetime.now().strftime("%d %b %Y")
    core = int((g.priority == "FILL FIRST").sum())

    a = df.groupby("crop_type").area_ha.sum()
    plant = a.get("PLANT", 0) + a.get("AUTUMN", 0)
    r1, r2 = a.get("RATOON", 0), a.get("RATOON II", 0)
    rec = df[df.crop_type != "RATOON"]
    top = df.groupby("variety").area_ha.sum().sort_values(ascending=False)
    measured = {
        "ratio": r1 / plant, "carry": r2 / r1 * 100,
        "lowland": rec[rec.land_type == "LOWLAND"].area_ha.sum() / rec.area_ha.sum() * 100,
        "top_share": top.iloc[0] / total * 100, "top_name": top.index[0],
        "total": total,
    }

    wb = Workbook()
    wb.remove(wb.active)
    sheet_start(wb, len(g), core, total, stamp)
    sheet_varieties(wb, g)
    sheet_rules(wb, measured)
    sheet_strategy(wb, g)
    sheet_reference(wb)
    wb.save(OUT)

    print(f"wrote {OUT}  ({os.path.getsize(OUT)/1024:.0f} KB)")
    print(f"  tabs        {', '.join(wb.sheetnames)}")
    print(f"  varieties   {len(g)}   FILL FIRST {core} ({g[g.priority=='FILL FIRST'].share.sum():.1f}% of area)")


if __name__ == "__main__":
    main()
