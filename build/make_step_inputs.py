# -*- coding: utf-8 -*-
"""Build the Step 3 and Step 4 input workbooks.

Step 1 is the ERP survey, so nothing is filled by hand there. Step 2 already has
its own workbook. These are the remaining two screens whose values a person sets:

    Step 3 - Agronomic Rules.xlsx     the mill's own settings, each shown beside
                                      what the survey actually measured
    Step 4 - Seed and Strategy.xlsx   where each variety should go, and how much
                                      cane to hold back as seed to get it there

Both follow the Step 2 template's shape: an input sheet, then a "How to fill
this" sheet. Measured columns come pre-filled; judgement columns are blank with
dropdowns so a typo cannot enter the engine.

    python build/make_step_inputs.py
"""

import datetime
import os
import subprocess
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CACHE = os.path.join(HERE, "cache", "plots_2627.parquet")
OUT_DIR = os.path.join(ROOT, "plan_inputs")

CORE_COVERAGE = 0.95
INK = "1F2937"
HDR_LOCKED = "334155"
HDR_FILL = "0F766E"
BAND_LOCKED = "F1F5F9"
BAND_CORE = "FEF9C3"
TEAL = "0F766E"
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STRATEGY = ["INTRODUCE-NEW", "EXPAND", "HOLD", "REDUCE", "EXIT"]
RETENTION = ["100% - all cane kept as seed", "50% - half kept", "25% - mostly crushed"]
BUD = ["DOUBLE BUD", "SINGLE BUD"]
RATOONS = ["1", "2"]


def banner(ws, ncol, title, subtitle, note):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = Font(size=10, italic=True, color="475569")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.cell(row=3, column=1, value=note).font = Font(size=9, color="94A3B8")


def header(ws, row, cols, locked_upto):
    for i, (label, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HDR_LOCKED if i <= locked_upto else HDR_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 34


def dropdown(ws, col, first, last, options, label):
    dv = DataValidation(
        type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
        showDropDown=False, errorTitle="Not an accepted value",
        error="Choose one of: " + ", ".join(options),
        promptTitle=label, prompt="Pick one of: " + ", ".join(options),
    )
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


def guard(ws, col, first, last, lo, hi, label):
    dv = DataValidation(
        type="decimal", operator="between", formula1=lo, formula2=hi, allow_blank=True,
        errorTitle="Out of range",
        error=f"{label} should be between {lo} and {hi}. Leave blank to keep the default.",
    )
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


def guide_sheet(wb, title, intro, sections, footer):
    ws = wb.create_sheet("How to fill this")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    r = [1]

    def line(a, b="", bold=True, size=10, color=INK, gap=False):
        if gap:
            r[0] += 1
        ws.cell(row=r[0], column=1, value=a).font = Font(bold=bold, size=size, color=color)
        c = ws.cell(row=r[0], column=2, value=b)
        c.font = Font(size=size, color=INK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0] += 1

    line(title, "", size=14, color=TEAL)
    line("", intro)
    for head, rows in sections:
        line(head, "", size=12, gap=True)
        for a, b in rows:
            line(a, b)
    line("", footer, gap=True)
    return ws


def load(df=None):
    if not os.path.exists(CACHE):
        subprocess.run([sys.executable, os.path.join(HERE, "extract_2627.py")], check=True)
    df = pd.read_parquet(CACHE)
    total = df.area_ha.sum()
    g = (df.groupby("variety").agg(area_ha=("area_ha", "sum")).reset_index())
    mat = df.groupby("variety").crop_category.agg(
        lambda x: x.mode().iat[0] if len(x.mode()) else "UNKNOWN")
    g["maturity"] = g.variety.map(mat).fillna("UNKNOWN")
    g = g[g.area_ha > 0].sort_values("area_ha", ascending=False).reset_index(drop=True)
    g["share"] = g.area_ha / total * 100
    g["cum"] = g.share.cumsum()
    g["priority"] = ["FILL FIRST" if c <= CORE_COVERAGE * 100 else "optional" for c in g.cum]
    return g, total, df


# ============================================================ STEP 3
def build_step3(measured, stamp):
    cols = [("Setting", 32), ("What it controls", 56), ("Current default", 15),
            ("Survey measured", 17), ("YOUR VALUE", 14), ("Notes", 42)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Agronomic Rules"
    banner(ws, len(cols),
           "STEP 3 - AGRONOMIC RULES",
           "The mill's own settings. Each shows what the 2026-27 survey actually measured beside it.",
           f"Leave YOUR VALUE blank to keep the current default.   Generated {stamp}.")
    HDR = 5
    header(ws, HDR, cols, 4)

    groups = [
        ("SEED AND MULTIPLICATION", [
            ("Seed rate (qtl/ha)", "Cane needed to plant one hectare. 65 qtl is 6.5 tonnes.",
             "65", "-", "", "num", 20, 140),
            ("Bud type", "SINGLE BUD uses about 30% less cane than DOUBLE BUD.",
             "DOUBLE BUD", "-", "", "bud", 0, 0),
            ("Multiplication factor", "1 ha of seed plants how many ha next year. It is really the ratio of nursery to farmer planting density, not a property of the cane.",
             "8", "-", "Confirm the two spacings", "num", 2, 20),
            ("Seed purchase ceiling (ha)", "Most that can be bought from the breeding institute in one year.",
             "35", "-", "", "num", 0, 500),
            ("Test plot size (ha)", "Size of a first trial of a new variety.",
             "5", "-", "", "num", 0, 100),
        ]),
        ("CROP CYCLE  -  decides how much land is free each year", [
            ("Ratoons taken", "How many ratoons before ploughing out. At 1, the carry rate below is ignored.",
             "2", "effectively 1", "Only 3.8% of ratoon fields reach a second ratoon", "ratoons", 0, 0),
            ("Ratoon : Plant ratio", "Ratoon hectares standing for every 1 hectare planted. This single number moves the replant budget by about 4,000 ha.",
             "0.90", f"{measured['ratio']:.2f}", "0.90 is the guiding document's target; 0.68 is what the survey finds", "num", 0.3, 2.0),
            ("Ratoon II carry rate (%)", "Share of ratoon I fields kept for a second ratoon.",
             "3.6", f"{measured['carry']:.1f}", "Default is already close to measured", "num", 0, 30),
        ]),
        ("SAFETY CAPS  -  manager owned", [
            ("Max variety concentration (%)", "No single variety above this share of the command area.",
             "40", f"{measured['top_share']:.1f} now", f"{measured['top_name']} is the largest", "num", 10, 100),
            ("Village concentration cap (%)", "No single village above this share on one variety. This is where a disease outbreak actually starts.",
             "60", "-", "", "num", 20, 100),
            ("Lowland coverage floor (%)", "Minimum share of the plan on lowland-capable varieties.",
             "28", f"{measured['lowland']:.1f}", "Measured is above the floor, so there is no breach", "num", 0, 80),
            ("Red rot trigger (%)", "Most area allowed on red-rot susceptible varieties before the plan is flagged.",
             "2", "-", "Needs the Red Rot column on the Step 2 sheet", "num", 0, 50),
        ]),
        ("MILL OUTPUT  -  decides what a plan is worth", [
            ("Juice to recovery factor", "Recovery as a share of juice sucrose. 18% sucrose x 0.635 gives 11.43% recovery.",
             "0.635", "-", "The mill's own conversion", "num", 0.4, 0.9),
            ("Season crush (MT cane)", "Cane crushed in a season, used to price a recovery change.",
             "1350000", "-", "", "num", 100000, 10000000),
            ("Sugar price (Rs/kg)", "Ex-mill price, used for the same purpose.",
             "38", "-", "", "num", 1, 200),
        ]),
        ("SCOPE", [
            ("Command area (ha)", "Total land the plan covers.",
             "57000", f"{measured['total']:,.0f}", "Survey is within 1% of the default", "num", 10000, 200000),
            ("Planning horizon (years)", "How many years the plan covers. 1 to 5.",
             "3", "-", "", "num", 1, 5),
            ("Base year", "The season the survey describes.",
             "2026-27", "-", "", "text", 0, 0),
        ]),
    ]

    r = HDR + 1
    for gname, rows in groups:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        c = ws.cell(row=r, column=1, value=gname)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="475569")
        ws.row_dimensions[r].height = 20
        r += 1
        for name, what, default, meas, note, kind, lo, hi in rows:
            for j, v in enumerate([name, what, default, meas, None, note], start=1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = BORDER
                c.font = Font(size=10, color=INK, bold=(j == 1))
                c.alignment = Alignment(wrap_text=(j in (2, 6)), vertical="top",
                                        horizontal="center" if j in (3, 4, 5) else "left")
                if j <= 4:
                    c.fill = PatternFill("solid", fgColor=BAND_LOCKED)
            if kind == "num":
                guard(ws, "E", r, r, lo, hi, name)
            elif kind == "bud":
                dropdown(ws, "E", r, r, BUD, name)
            elif kind == "ratoons":
                dropdown(ws, "E", r, r, RATOONS, name)
            ws.row_dimensions[r].height = 32
            r += 1

    ws.freeze_panes = ws.cell(row=HDR + 1, column=1)

    guide_sheet(
        wb, "STEP 3 - HOW TO FILL THIS",
        "These are the mill's own planning rules - not facts from the survey, but decisions. Every one "
        "of them changes the plan, so each is shown beside what the survey actually measured. Set only "
        "what you want to change; a blank YOUR VALUE keeps the current default.",
        [
            ("THE TWO THAT MATTER MOST", [
                ("Ratoon : Plant ratio",
                 "The default of 0.90 comes from the guiding document - it is the target. The survey measures 0.68 - "
                 "that is what happens. The gap is worth about 4,000 hectares of replant budget, and every figure "
                 "downstream scales with it. Neither is wrong; the engine just has to be told which one to plan against."),
                ("Ratoons taken",
                 "Set to 2 by default, but only 3.8% of ratoon fields at Gobind reach a second ratoon. At 2, a field "
                 "planted in Year 1 does not come free again until Year 4, so the third year of the plan has almost "
                 "nothing to plant. At 1, it comes back in Year 3 and the plan covers about 39% more land."),
            ]),
            ("WHAT EACH GROUP DOES", [
                ("Seed and multiplication", "Converts hectares into quintals of cane, and decides how fast a new variety can arrive. "
                                            "At 35 ha bought and 8x multiplication: 35 -> 280 -> 2,240 ha in two cycles."),
                ("Crop cycle", "Decides how much land is free to replant each year. At Gobind only about 23,300 ha of 56,491 is free in any year."),
                ("Safety caps", "The rules a plan must not break. Concentration, village concentration, lowland cover, red rot exposure."),
                ("Mill output", "Turns a varietal mix into recovery, and recovery into rupees."),
                ("Scope", "How big the command area is and how many years to plan for."),
            ]),
            ("A NOTE ON THE MEASURED COLUMN", [
                ("Where it comes from", "The 2026-27 plot survey, computed the same way the app computes it. A dash means the survey has "
                                        "nothing to say about that setting - it is purely a mill decision."),
                ("Lowland floor", "Reads 39.5%, comfortably above the 28% floor. An earlier figure of 23.8% was wrong: it counted ratoon "
                                  "plots as upland, and the ERP records no land type on those at all."),
            ]),
        ],
        "When you are done, save the file and send it back. Do not rename the columns or move rows between "
        "groups - the system reads them by name.",
    )
    path = os.path.join(OUT_DIR, "Step 3 - Agronomic Rules.xlsx")
    wb.save(path)
    return path, sum(len(rows) for _, rows in groups)


# ============================================================ STEP 4
def build_step4(g, stamp):
    cols = [("#", 5), ("Variety", 22), ("Area (ha)", 12), ("% of Area", 10),
            ("Maturity", 12), ("Priority", 12),
            ("Strategy", 20), ("Seed Retention", 28),
            ("Target Y1 (ha)", 14), ("Target Y2 (ha)", 14), ("Target Y3 (ha)", 14),
            ("Notes", 40)]
    LOCKED = 6
    wb = Workbook()
    ws = wb.active
    ws.title = "Seed and Strategy"
    banner(ws, len(cols),
           "STEP 4 - SEED AND STRATEGY",
           "Where each variety should go, and how much cane to hold back as seed to get it there.",
           f"Grey columns come from the survey. Targets are optional.   Generated {stamp}.")
    HDR = 5
    header(ws, HDR, cols, LOCKED)
    first = HDR + 1

    for i, row in enumerate(g.itertuples()):
        r = first + i
        vals = [i + 1, row.variety, round(float(row.area_ha), 1),
                round(float(row.share), 2), row.maturity, row.priority]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = Font(size=10, color=INK, bold=(j == 2))
            fill = BAND_CORE if (j == 6 and row.priority == "FILL FIRST") else BAND_LOCKED
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(horizontal="right" if j in (1, 3, 4) else "left")
        for j in range(LOCKED + 1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.border = BORDER
            c.font = Font(size=10, color=INK)

    last = first + len(g) - 1
    dropdown(ws, "G", first, last, STRATEGY, "Strategy")
    dropdown(ws, "H", first, last, RETENTION, "Seed Retention")
    for col in ("I", "J", "K"):
        guard(ws, col, first, last, 0, 60000, "Target hectares")

    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{last}"

    core = int((g.priority == "FILL FIRST").sum())
    guide_sheet(
        wb, "STEP 4 - HOW TO FILL THIS",
        "Step 2 said what each variety IS. This says where it should GO. Two columns matter - Strategy and "
        "Seed Retention. The three target columns are optional and most people leave them blank.",
        [
            ("THE ONE IDEA BEHIND THIS SHEET", [
                ("Seed is cane",
                 "There is no separate seed packet. To grow more of a variety you cut up cane and plant the pieces. "
                 "So every hectare is either crushed for sugar, or held back as seed for next year. You cannot do both "
                 "with the same cane. Seed Retention is that choice, and it is the real trade-off in the whole plan."),
            ]),
            ("STRATEGY  -  what you want the variety to do", [
                ("EXPAND", "Grow it. How far it can actually get depends on Seed Retention and its own multiplication."),
                ("HOLD", "Keep it roughly where it is."),
                ("REDUCE", "Bring it down gradually, about a quarter each year."),
                ("EXIT", "Stop it. Its plots are freed first, because those are the ones the plan most wants to move."),
                ("INTRODUCE-NEW", "Start it from a nursery. Only for a variety farmers do not already grow."),
            ]),
            ("SEED RETENTION  -  how much cane to hold back", [
                ("100%", "Everything becomes seed. Fastest growth, but no sugar from that variety this year."),
                ("50%", "Half and half. The usual balance."),
                ("25%", "Mostly crushed. Right for a variety that has finished growing."),
                ("What it is worth", "A variety standing on 2,062 ha with 8x multiplication reaches about 2,784 ha at 25% retention, "
                                     "4,227 ha at 50%, and 7,114 ha at 100%. That is the price of expansion, in sugar forgone."),
            ]),
            ("TARGET Y1 / Y2 / Y3  -  optional", [
                ("Leave blank normally", "The system works out how far seed can take each variety on its own. That is the realistic answer."),
                ("Fill only if", "Somebody senior has asked for a specific hectarage by a specific year. Then the system will tell you "
                                 "what is missing - and if the target cannot be reached, it says so rather than quietly trimming it."),
            ]),
            ("WHERE TO START", [
                ("The FILL FIRST rows", f"There are {core} of them and they cover 95% of the command area. The other {len(g) - core} "
                                        "varieties average about 44 ha each and can wait."),
                ("A variety you are unsure about", "Leave Strategy blank rather than guessing. A blank the system asks about is far better "
                                                   "than a direction it acts on."),
            ]),
        ],
        "When you are done, save the file and send it back. Do not rename the columns or reorder them - "
        "the system reads them by name.",
    )
    path = os.path.join(OUT_DIR, "Step 4 - Seed and Strategy.xlsx")
    wb.save(path)
    return path, core


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    g, total, df = load()
    stamp = datetime.datetime.now().strftime("%d %b %Y")

    a = df.groupby("crop_type").area_ha.sum()
    plant = a.get("PLANT", 0) + a.get("AUTUMN", 0)
    r1, r2 = a.get("RATOON", 0), a.get("RATOON II", 0)
    rec = df[df.crop_type != "RATOON"]
    top = df.groupby("variety").area_ha.sum().sort_values(ascending=False)
    measured = {
        "ratio": r1 / plant, "carry": r2 / r1 * 100,
        "lowland": rec[rec.land_type == "LOWLAND"].area_ha.sum() / rec.area_ha.sum() * 100,
        "top_share": top.iloc[0] / total * 100, "top_name": top.index[0], "total": total,
    }

    p3, n3 = build_step3(measured, stamp)
    p4, core = build_step4(g, stamp)
    for p, note in [(p3, f"{n3} settings"), (p4, f"{len(g)} varieties, {core} marked FILL FIRST")]:
        print(f"wrote {p}")
        print(f"      {os.path.getsize(p)/1024:.0f} KB  -  {note}")


if __name__ == "__main__":
    main()
